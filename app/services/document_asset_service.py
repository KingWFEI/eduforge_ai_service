import json
import re
import shutil
import zipfile
from dataclasses import asdict, dataclass
from pathlib import Path


_MARKDOWN_IMAGE_RE = re.compile(r"!\[[^\]]*]\((?P<url>[^)]+)\)")
_SLIDE_MARKER_RE = re.compile(r"(<!--\s*Slide number:\s*(?P<number>\d+)\s*-->)")
_PAGE_MARKER_RE = re.compile(r"(<!--\s*Page number:\s*(?P<number>\d+)\s*-->)")


@dataclass(frozen=True)
class DocumentAsset:
    filename: str
    url: str
    asset_type: str
    page_no: int | None = None
    slide_no: int | None = None
    alt: str = "文档插图"


def prepare_document_assets(
    markdown: str,
    file_path: str,
    course_id: str,
    document_id: str,
) -> tuple[str, list[DocumentAsset]]:
    source = Path(file_path)
    assets_dir = (
        Path("uploads")
        / "course_documents"
        / course_id
        / document_id
        / "assets"
    )
    if assets_dir.exists():
        shutil.rmtree(assets_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)

    suffix = source.suffix.lower()
    if suffix == ".pptx":
        assets = _extract_pptx_images(source, assets_dir, course_id, document_id)
        enriched = _inject_pptx_images(markdown, assets)
    elif suffix == ".pdf":
        assets = _render_pdf_image_pages(source, assets_dir, course_id, document_id)
        enriched = _inject_pdf_page_images(markdown, assets)
    elif suffix == ".docx":
        assets = _extract_docx_media(source, assets_dir, course_id, document_id)
        enriched = _append_asset_gallery(markdown, assets, "文档原始插图")
    else:
        assets = []
        enriched = markdown

    _write_manifest(assets_dir.parent, assets)
    return enriched, assets


def load_document_assets(course_id: str, document_id: str) -> list[dict]:
    manifest = (
        Path("uploads")
        / "course_documents"
        / course_id
        / document_id
        / "manifest.json"
    )
    if not manifest.exists():
        return []
    try:
        data = json.loads(manifest.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return []
    return data if isinstance(data, list) else []


def extract_image_urls(markdown: str | None) -> list[str]:
    if not markdown:
        return []
    urls = []
    for match in _MARKDOWN_IMAGE_RE.finditer(markdown):
        url = match.group("url").strip()
        if url.startswith("/uploads/") and url not in urls:
            urls.append(url)
    return urls


def append_source_images(content_markdown: str | None, source_chunks: list[dict]) -> str | None:
    if not content_markdown:
        return content_markdown
    urls: list[str] = []
    for chunk in source_chunks:
        for url in extract_image_urls(str(chunk.get("content") or "")):
            if url not in urls:
                urls.append(url)
    if not urls:
        return content_markdown

    gallery = ["", "## 来源插图", ""]
    gallery.extend(f"![来源插图]({url})" for url in urls[:12])
    return content_markdown.rstrip() + "\n" + "\n\n".join(gallery)


def _extract_pptx_images(
    source: Path,
    assets_dir: Path,
    course_id: str,
    document_id: str,
) -> list[DocumentAsset]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    assets: list[DocumentAsset] = []
    presentation = Presentation(str(source))
    for slide_no, slide in enumerate(presentation.slides, start=1):
        image_no = 0
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image_no += 1
            image = shape.image
            extension = (image.ext or "png").lower()
            filename = f"slide-{slide_no:03d}-image-{image_no:02d}.{extension}"
            (assets_dir / filename).write_bytes(image.blob)
            assets.append(
                DocumentAsset(
                    filename=filename,
                    url=_asset_url(course_id, document_id, filename),
                    asset_type="slide_image",
                    slide_no=slide_no,
                    alt=f"第 {slide_no} 页插图 {image_no}",
                )
            )
    return assets


def _render_pdf_image_pages(
    source: Path,
    assets_dir: Path,
    course_id: str,
    document_id: str,
) -> list[DocumentAsset]:
    import pdfplumber
    import pypdfium2 as pdfium

    image_pages: list[int] = []
    with pdfplumber.open(str(source)) as pdf:
        for page_no, page in enumerate(pdf.pages, start=1):
            if page.images:
                image_pages.append(page_no)

    assets: list[DocumentAsset] = []
    pdf = pdfium.PdfDocument(str(source))
    try:
        for page_no in image_pages:
            filename = f"page-{page_no:03d}.png"
            page = pdf[page_no - 1]
            bitmap = page.render(scale=1.5)
            bitmap.to_pil().save(assets_dir / filename, format="PNG")
            page.close()
            assets.append(
                DocumentAsset(
                    filename=filename,
                    url=_asset_url(course_id, document_id, filename),
                    asset_type="pdf_page",
                    page_no=page_no,
                    alt=f"PDF 第 {page_no} 页",
                )
            )
    finally:
        pdf.close()
    return assets


def _extract_docx_media(
    source: Path,
    assets_dir: Path,
    course_id: str,
    document_id: str,
) -> list[DocumentAsset]:
    assets: list[DocumentAsset] = []
    with zipfile.ZipFile(source) as archive:
        media_names = sorted(
            name for name in archive.namelist() if name.startswith("word/media/")
        )
        for index, media_name in enumerate(media_names, start=1):
            extension = Path(media_name).suffix.lower() or ".bin"
            filename = f"image-{index:03d}{extension}"
            (assets_dir / filename).write_bytes(archive.read(media_name))
            assets.append(
                DocumentAsset(
                    filename=filename,
                    url=_asset_url(course_id, document_id, filename),
                    asset_type="document_image",
                    alt=f"文档插图 {index}",
                )
            )
    return assets


def _inject_pptx_images(markdown: str, assets: list[DocumentAsset]) -> str:
    assets_by_slide: dict[int, list[DocumentAsset]] = {}
    for asset in assets:
        if asset.slide_no is not None:
            assets_by_slide.setdefault(asset.slide_no, []).append(asset)

    markdown = re.sub(r"(?m)^\s*!\[[^\]]*]\([^)]+\)\s*$", "", markdown)

    def replacement(match: re.Match) -> str:
        slide_no = int(match.group("number"))
        refs = assets_by_slide.get(slide_no, [])
        if not refs:
            return match.group(1)
        images = "\n\n".join(f"![{asset.alt}]({asset.url})" for asset in refs)
        return f"{match.group(1)}\n\n{images}"

    return _SLIDE_MARKER_RE.sub(replacement, markdown)


def _inject_pdf_page_images(markdown: str, assets: list[DocumentAsset]) -> str:
    assets_by_page = {
        asset.page_no: asset
        for asset in assets
        if asset.page_no is not None
    }

    def replacement(match: re.Match) -> str:
        page_no = int(match.group("number"))
        asset = assets_by_page.get(page_no)
        if asset is None:
            return match.group(1)
        return f"{match.group(1)}\n\n![{asset.alt}]({asset.url})"

    return _PAGE_MARKER_RE.sub(replacement, markdown)


def _append_asset_gallery(
    markdown: str,
    assets: list[DocumentAsset],
    title: str,
) -> str:
    if not assets:
        return markdown
    gallery = [f"## {title}"]
    gallery.extend(f"![{asset.alt}]({asset.url})" for asset in assets)
    return markdown.rstrip() + "\n\n" + "\n\n".join(gallery)


def _write_manifest(document_dir: Path, assets: list[DocumentAsset]) -> None:
    document_dir.mkdir(parents=True, exist_ok=True)
    (document_dir / "manifest.json").write_text(
        json.dumps([asdict(asset) for asset in assets], ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def _asset_url(course_id: str, document_id: str, filename: str) -> str:
    return (
        f"/uploads/course_documents/{course_id}/{document_id}/assets/{filename}"
    )

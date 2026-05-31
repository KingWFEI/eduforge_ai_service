import json
import os
import uuid
from pathlib import Path
from typing import Any, List

import chromadb
import numpy as np
from docx import Document
from fastapi import UploadFile, status
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.services.llm_service import LLMService
from app.utils.response import AppException, ErrorCode

_MODEL = None
_CHROMA_CLIENT = None
_COLLECTION = None

CHROMA_DIR = "chroma_db"
COLLECTION_NAME = "eduforge_knowledge_chunks"
EMBEDDING_MODEL_NAME = os.getenv(
    "EMBEDDING_MODEL_NAME",
    "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2",
)
EMBEDDING_LOCAL_FILES_ONLY = os.getenv("EMBEDDING_LOCAL_FILES_ONLY", "false").lower() == "true"
SUPPORTED_DOCUMENT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".py",
    ".dot",
    ".pdf",
    ".docx",
    ".ipynb",
}
SUPPORTED_DOCUMENT_TYPES_MESSAGE = "仅支持 txt、md、csv、json、py、dot、pdf、docx、ipynb 文件"


def get_embedding_model():
    """
    懒加载向量模型。
    第一次调用会下载模型，后面复用。
    """
    global _MODEL

    if _MODEL is None:
        try:
            _MODEL = SentenceTransformer(
                EMBEDDING_MODEL_NAME,
                local_files_only=EMBEDDING_LOCAL_FILES_ONLY,
            )
        except Exception as exc:
            raise AppException(
                code=ErrorCode.KNOWLEDGE_RETRIEVAL_ERROR,
                message=(
                    "向量模型未就绪，暂时无法解析并索引课程资料。"
                    "请检查服务器网络，或提前下载/缓存 embedding 模型后重试。"
                ),
                status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            ) from exc

    return _MODEL


def ensure_embedding_model_ready() -> None:
    get_embedding_model()


def embed_text(text_value: str) -> List[float]:
    model = get_embedding_model()
    vector = model.encode(text_value, normalize_embeddings=True)
    return vector.tolist()


def get_chroma_collection():
    """
    获取 Chroma collection。
    所有课程的知识块都放在同一个 collection 里，
    通过 metadata.course_id 过滤。
    """
    global _CHROMA_CLIENT, _COLLECTION

    if _CHROMA_CLIENT is None:
        _CHROMA_CLIENT = chromadb.PersistentClient(path=CHROMA_DIR)

    if _COLLECTION is None:
        _COLLECTION = _CHROMA_CLIENT.get_or_create_collection(
            name=COLLECTION_NAME,
            metadata={"hnsw:space": "cosine"},
        )

    return _COLLECTION


def extract_text_from_file(file_path: str) -> str:
    """
    从上传文件中提取文本。
    支持：txt、md、csv、json、py、dot、pdf、docx、ipynb
    """
    suffix = Path(file_path).suffix.lower()

    if suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise AppException(
            code=ErrorCode.PARAM_ERROR,
            message=SUPPORTED_DOCUMENT_TYPES_MESSAGE,
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    if suffix in [".txt", ".md", ".py", ".json", ".csv", ".dot"]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    if suffix == ".pdf":
        reader = PdfReader(file_path)
        texts = []

        for page in reader.pages:
            page_text = page.extract_text() or ""
            texts.append(page_text)

        return "\n".join(texts)

    if suffix == ".docx":
        doc = Document(file_path)
        return "\n".join([p.text for p in doc.paragraphs])

    if suffix == ".ipynb":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            notebook = json.load(f)

        texts = []

        for cell in notebook.get("cells", []):
            cell_type = cell.get("cell_type")
            source = cell.get("source", [])

            if isinstance(source, list):
                source_text = "".join(source)
            else:
                source_text = str(source)

            if not source_text.strip():
                continue

            if cell_type == "markdown":
                texts.append("【Markdown说明】\n" + source_text)
            elif cell_type == "code":
                texts.append("【代码单元】\n" + source_text)
            else:
                texts.append(source_text)

        return "\n\n".join(texts)

    raise AppException(
        code=ErrorCode.PARAM_ERROR,
        message=SUPPORTED_DOCUMENT_TYPES_MESSAGE,
        status_code=status.HTTP_400_BAD_REQUEST,
    )


def split_text(text_value: str, chunk_size: int = 500, overlap: int = 80) -> List[str]:
    """
    简单按字符切块。
    先跑通 RAG，后面再优化为按标题、段落切分。
    """
    text_value = text_value.replace("\r\n", "\n").strip()

    if not text_value:
        return []

    chunks = []
    start = 0

    while start < len(text_value):
        end = start + chunk_size
        chunk = text_value[start:end].strip()

        if chunk:
            chunks.append(chunk)

        start = end - overlap

        if start < 0:
            start = 0

        if start >= len(text_value):
            break

    return chunks


def upload_and_index_course_document(
    db: Session,
    course_id: str,
    file: UploadFile,
    current_user_id: str | None = None,
    chapter_id: str | None = None,
    description: str | None = None,
) -> dict:
    """
    上传课程资料：
    1. 保存文件到本地
    2. 保存文档记录到 course_documents
    3. 创建 vector_index_records 索引任务记录
    4. 解析文本
    5. 切分 chunks
    6. chunks 写入 Chroma
    7. chunks 同步写入 MySQL knowledge_chunks
    8. 更新索引任务状态
    """

    if not file.filename:
        raise AppException(
            code=ErrorCode.PARAM_ERROR,
            message="文件名不能为空",
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    file_suffix = Path(file.filename).suffix.lower()

    if file_suffix not in SUPPORTED_DOCUMENT_EXTENSIONS:
        raise AppException(
            code=ErrorCode.PARAM_ERROR,
            message=SUPPORTED_DOCUMENT_TYPES_MESSAGE,
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    # 1. 确认课程存在
    course = db.execute(
        text(
            """
            SELECT course_id
            FROM courses
            WHERE course_id = :course_id
            LIMIT 1
            """
        ),
        {"course_id": course_id},
    ).mappings().first()

    if course is None:
        raise AppException(
            code=ErrorCode.NOT_FOUND,
            message="课程不存在，请先创建课程后再上传资料",
            status_code=status.HTTP_404_NOT_FOUND,
        )

    if chapter_id:
        chapter = db.execute(
            text(
                """
                SELECT id
                FROM course_chapters
                WHERE id = :chapter_id
                  AND course_id = :course_id
                LIMIT 1
                """
            ),
            {
                "chapter_id": chapter_id,
                "course_id": course_id,
            },
        ).mappings().first()

        if chapter is None:
            raise AppException(
                code=ErrorCode.NOT_FOUND,
                message="章节不存在或不属于该课程",
                status_code=status.HTTP_404_NOT_FOUND,
            )

    ensure_embedding_model_ready()

    document_id = "doc_" + uuid.uuid4().hex[:12]
    index_record_id = "idx_" + uuid.uuid4().hex[:12]

    upload_dir = Path("uploads") / "course_documents" / course_id
    upload_dir.mkdir(parents=True, exist_ok=True)

    safe_filename = file.filename.replace("\\", "_").replace("/", "_")
    file_path = upload_dir / f"{document_id}_{safe_filename}"

    file_bytes = file.file.read()

    with open(file_path, "wb") as f:
        f.write(file_bytes)

    file_type = file_suffix.replace(".", "")
    file_size = len(file_bytes)

    try:
        # 2. 先插入文档记录，初始状态为 processing
        db.execute(
            text(
                """
                INSERT INTO course_documents (
                    id,
                    course_id,
                    chapter_id,
                    filename,
                    file_path,
                    file_type,
                    file_size,
                    description,
                    status,
                    parse_status,
                    index_status,
                    chunk_count,
                    uploaded_by,
                    uploaded_at,
                    updated_at
                )
                VALUES (
                    :id,
                    :course_id,
                    :chapter_id,
                    :filename,
                    :file_path,
                    :file_type,
                    :file_size,
                    :description,
                    'active',
                    'processing',
                    'processing',
                    0,
                    :uploaded_by,
                    NOW(),
                    NOW()
                )
                """
            ),
            {
                "id": document_id,
                "course_id": course_id,
                "chapter_id": chapter_id,
                "filename": file.filename,
                "file_path": str(file_path),
                "file_type": file_type,
                "file_size": file_size,
                "description": description,
                "uploaded_by": current_user_id,
            },
        )

        # 3. 创建索引任务记录
        db.execute(
            text(
                """
                INSERT INTO vector_index_records (
                    id,
                    course_id,
                    document_id,
                    index_type,
                    collection_name,
                    status,
                    chunk_count,
                    success_count,
                    failed_count,
                    error_message,
                    started_at,
                    finished_at,
                    created_by,
                    created_at
                )
                VALUES (
                    :id,
                    :course_id,
                    :document_id,
                    'chroma',
                    :collection_name,
                    'processing',
                    0,
                    0,
                    0,
                    NULL,
                    NOW(),
                    NULL,
                    :created_by,
                    NOW()
                )
                """
            ),
            {
                "id": index_record_id,
                "course_id": course_id,
                "document_id": document_id,
                "collection_name": COLLECTION_NAME,
                "created_by": current_user_id,
            },
        )

        db.commit()

        # 4. 解析文本
        raw_text = extract_text_from_file(str(file_path))
        chunks = split_text(raw_text)

        if not chunks:
            raise AppException(
                code=ErrorCode.PARAM_ERROR,
                message="文件内容为空，无法生成知识块。如果是扫描版 PDF，需要 OCR。",
                status_code=status.HTTP_400_BAD_REQUEST,
            )

        # 5. 写入 Chroma，同时准备写 MySQL knowledge_chunks
        collection = get_chroma_collection()

        ids = []
        documents = []
        embeddings = []
        metadatas = []

        mysql_chunk_rows = []

        for index, chunk in enumerate(chunks):
            chunk_id = "chunk_" + uuid.uuid4().hex[:12]
            section = f"{file.filename} - 片段 {index + 1}"
            embedding = embed_text(chunk)

            # Chroma 数据
            ids.append(chunk_id)
            documents.append(chunk)
            embeddings.append(embedding)
            metadata = {
                "course_id": course_id,
                "document_id": document_id,
                "filename": file.filename,
                "chunk_index": index,
                "section": section,
                "vector_id": chunk_id,
            }
            if chapter_id:
                metadata["chapter_id"] = chapter_id
            metadatas.append(metadata)

            # MySQL knowledge_chunks 数据
            mysql_chunk_rows.append(
                {
                    "id": chunk_id,
                    "course_id": course_id,
                    "document_id": document_id,
                    "chapter_id": chapter_id,
                    "section": section,
                    "content": chunk,
                    "keywords_json": json.dumps([], ensure_ascii=False),
                    "chunk_index": index,
                    "vector_id": chunk_id,
                }
            )

        collection.add(
            ids=ids,
            documents=documents,
            embeddings=embeddings,
            metadatas=metadatas,
        )

        # 6. 同步写入 MySQL knowledge_chunks
        for row in mysql_chunk_rows:
            db.execute(
                text(
                    """
                    INSERT INTO knowledge_chunks (
                        id,
                        course_id,
                        document_id,
                        chapter_id,
                        knowledge_point_id,
                        section,
                        content,
                        keywords_json,
                        page_no,
                        chunk_index,
                        vector_id,
                        indexed,
                        deleted,
                        created_at
                    )
                    VALUES (
                        :id,
                        :course_id,
                        :document_id,
                        :chapter_id,
                        NULL,
                        :section,
                        :content,
                        :keywords_json,
                        NULL,
                        :chunk_index,
                        :vector_id,
                        1,
                        FALSE,
                        NOW()
                    )
                    """
                ),
                row,
            )

        # 7. 更新 course_documents 状态
        db.execute(
            text(
                """
                UPDATE course_documents
                SET parse_status = 'parsed',
                    index_status = 'indexed',
                    chunk_count = :chunk_count,
                    updated_at = NOW()
                WHERE id = :document_id
                  AND course_id = :course_id
                """
            ),
            {
                "chunk_count": len(chunks),
                "document_id": document_id,
                "course_id": course_id,
            },
        )

        # 8. 更新索引任务记录
        db.execute(
            text(
                """
                UPDATE vector_index_records
                SET status = 'completed',
                    chunk_count = :chunk_count,
                    success_count = :success_count,
                    failed_count = 0,
                    error_message = NULL,
                    finished_at = NOW()
                WHERE id = :index_record_id
                """
            ),
            {
                "chunk_count": len(chunks),
                "success_count": len(chunks),
                "index_record_id": index_record_id,
            },
        )

        db.commit()

        return {
            "document_id": document_id,
            "course_id": course_id,
            "filename": file.filename,
            "chunk_count": len(chunks),
            "parse_status": "parsed",
            "index_status": "indexed",
        }

    except Exception as e:
        db.rollback()

        error_message = str(e)

        # 尽量更新文档失败状态
        try:
            db.execute(
                text(
                    """
                    UPDATE course_documents
                    SET parse_status = 'failed',
                        index_status = 'failed',
                        updated_at = NOW()
                    WHERE id = :document_id
                      AND course_id = :course_id
                    """
                ),
                {
                    "document_id": document_id,
                    "course_id": course_id,
                },
            )

            db.execute(
                text(
                    """
                    UPDATE vector_index_records
                    SET status = 'failed',
                        error_message = :error_message,
                        finished_at = NOW()
                    WHERE id = :index_record_id
                    """
                ),
                {
                    "index_record_id": index_record_id,
                    "error_message": error_message,
                },
            )

            db.commit()
        except Exception:
            db.rollback()

        if isinstance(e, AppException):
            raise e

        raise AppException(
            code=ErrorCode.SERVER_ERROR,
            message=f"资料解析或 Chroma 向量索引失败：{error_message}",
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
        )


def build_chroma_where(
    course_id: str,
    chapter_id: str | None = None,
    knowledge_point_id: str | None = None,
) -> dict:
    """
    构造 Chroma metadata 过滤条件。

    只传 course_id：
    where = {"course_id": "course_ai_basic"}

    传 course_id + chapter_id：
    where = {
        "$and": [
            {"course_id": "course_ai_basic"},
            {"chapter_id": "ch_xxx"}
        ]
    }
    """

    conditions = [
        {"course_id": course_id}
    ]

    if chapter_id:
        conditions.append({"chapter_id": chapter_id})

    if knowledge_point_id:
        conditions.append({"knowledge_point_id": knowledge_point_id})

    if len(conditions) == 1:
        return conditions[0]

    return {
        "$and": conditions
    }


def search_knowledge_chunks(
    db: Session,
    course_id: str,
    query: str,
    top_k: int = 5,
    chapter_id: str | None = None,
    knowledge_point_id: str | None = None,
) -> dict:
    """
    在 Chroma 中做向量检索。

    支持：
    1. 按课程检索
    2. 按课程 + 章节检索
    3. 按课程 + 章节 + 知识点检索
    """

    if not query.strip():
        raise AppException(
            code=ErrorCode.PARAM_ERROR,
            message="检索问题不能为空",
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    # 1. 先确认课程存在
    course = db.execute(
        text(
            """
            SELECT course_id
            FROM courses
            WHERE course_id = :course_id
            LIMIT 1
            """
        ),
        {"course_id": course_id},
    ).mappings().first()

    if course is None:
        raise AppException(
            code=ErrorCode.NOT_FOUND,
            message="课程不存在",
            status_code=status.HTTP_404_NOT_FOUND,
        )

    # 2. 如果传了 chapter_id，确认章节存在
    if chapter_id:
        chapter = db.execute(
            text(
                """
                SELECT id
                FROM course_chapters
                WHERE id = :chapter_id
                  AND course_id = :course_id
                LIMIT 1
                """
            ),
            {
                "chapter_id": chapter_id,
                "course_id": course_id,
            },
        ).mappings().first()

        if chapter is None:
            raise AppException(
                code=ErrorCode.NOT_FOUND,
                message="章节不存在或不属于该课程",
                status_code=status.HTTP_404_NOT_FOUND,
            )

    # 3. 如果传了 knowledge_point_id，确认知识点存在
    if knowledge_point_id:
        knowledge_point = db.execute(
            text(
                """
                SELECT id
                FROM knowledge_points
                WHERE id = :knowledge_point_id
                  AND course_id = :course_id
                LIMIT 1
                """
            ),
            {
                "knowledge_point_id": knowledge_point_id,
                "course_id": course_id,
            },
        ).mappings().first()

        if knowledge_point is None:
            raise AppException(
                code=ErrorCode.NOT_FOUND,
                message="知识点不存在或不属于该课程",
                status_code=status.HTTP_404_NOT_FOUND,
            )

    collection = get_chroma_collection()

    query_embedding = embed_text(query)

    where_filter = build_chroma_where(
        course_id=course_id,
        chapter_id=chapter_id,
        knowledge_point_id=knowledge_point_id,
    )

    result = collection.query(
        query_embeddings=[query_embedding],
        n_results=top_k,
        where=where_filter,
        include=["documents", "metadatas", "distances"],
    )

    ids = result.get("ids", [[]])[0]
    documents = result.get("documents", [[]])[0]
    metadatas = result.get("metadatas", [[]])[0]
    distances = result.get("distances", [[]])[0]

    active_chunk_ids = set()
    if ids:
        id_params = {f"id_{index}": chunk_id for index, chunk_id in enumerate(ids)}
        placeholders = ", ".join(f":{key}" for key in id_params)
        active_rows = db.execute(
            text(
                f"""
                SELECT kc.id
                FROM knowledge_chunks kc
                JOIN course_documents cd ON cd.id = kc.document_id
                WHERE kc.id IN ({placeholders})
                  AND COALESCE(kc.deleted, 0) = 0
                  AND COALESCE(cd.status, 'active') <> 'deleted'
                """
            ),
            id_params,
        ).mappings().all()
        active_chunk_ids = {row["id"] for row in active_rows}

    items = []

    for index, chunk_id in enumerate(ids):
        if chunk_id not in active_chunk_ids:
            continue

        metadata = metadatas[index] or {}
        distance = distances[index]

        # cosine 距离越小越相似，转成越大越好的 score
        score = round(1 - float(distance), 4)

        items.append(
            {
                "chunk_id": chunk_id,
                "document_id": metadata.get("document_id"),
                "course_id": metadata.get("course_id"),
                "chapter_id": metadata.get("chapter_id"),
                "knowledge_point_id": metadata.get("knowledge_point_id"),
                "section": metadata.get("section"),
                "content": documents[index],
                "chunk_index": metadata.get("chunk_index", 0),
                "score": score,
            }
        )

    return {
        "course_id": course_id,
        "query": query,
        "total": len(items),
        "items": items,
    }


def list_knowledge_chunks(
    db: Session,
    course_id: str,
) -> dict:
    """
    查看某门课程在 Chroma 中的知识块。
    """

    collection = get_chroma_collection()

    result = collection.get(
        where={"course_id": course_id},
        include=["documents", "metadatas"],
    )

    ids = result.get("ids", [])
    documents = result.get("documents", [])
    metadatas = result.get("metadatas", [])

    items = []

    for index, chunk_id in enumerate(ids):
        metadata = metadatas[index] or {}

        items.append(
            {
                "chunk_id": chunk_id,
                "document_id": metadata.get("document_id"),
                "course_id": metadata.get("course_id"),
                "section": metadata.get("section"),
                "content": documents[index],
                "chunk_index": metadata.get("chunk_index", 0),
                "indexed": 1,
            }
        )

    items.sort(key=lambda item: (item["document_id"] or "", item["chunk_index"]))

    return {
        "course_id": course_id,
        "total": len(items),
        "items": items,
    }


def delete_document_chunks_from_chroma(document_id: str) -> int:
    """
    根据 document_id 删除 Chroma 中的知识块。
    """

    collection = get_chroma_collection()

    existing = collection.get(
        where={"document_id": document_id},
        include=["metadatas"],
    )

    ids = existing.get("ids", [])

    if not ids:
        return 0

    collection.delete(ids=ids)

    return len(ids)


def update_document_chunks_metadata_in_chroma(
    document_id: str,
    metadata_patch: dict,
) -> int:
    """
    更新 Chroma 中某个文档下所有 chunks 的 metadata。

    用途：
    当课程资料关联章节 / 知识点后，同步更新 Chroma metadata，
    这样后续可以按 chapter_id / knowledge_point_id 过滤检索。
    """

    collection = get_chroma_collection()

    existing = collection.get(
        where={"document_id": document_id},
        include=["metadatas"],
    )

    ids = existing.get("ids", [])
    metadatas = existing.get("metadatas", [])

    if not ids:
        return 0

    new_metadatas = []

    for old_metadata in metadatas:
        merged = dict(old_metadata or {})

        for key, value in metadata_patch.items():
            if value is not None:
                merged[key] = value

        new_metadatas.append(merged)

    collection.update(
        ids=ids,
        metadatas=new_metadatas,
    )

    return len(ids)


def _build_rag_context(search_items: list[dict]) -> str:
    """
    把检索到的 chunks 拼成给大模型看的上下文。
    """

    context_parts = []

    for index, item in enumerate(search_items, start=1):
        section = item.get("section") or "未知来源"
        content = item.get("content") or ""

        context_parts.append(
            f"【资料片段 {index}】\n"
            f"来源：{section}\n"
            f"内容：\n{content}\n"
        )

    return "\n".join(context_parts)


def _build_rag_system_prompt() -> str:
    """
    RAG 问答系统提示词。
    """

    return (
        "你是 EduForge-AI 智学工坊的课程知识库问答助手。\n"
        "你必须严格基于用户提供的【资料片段】回答问题。\n"
        "如果资料片段中没有足够信息，请明确说明：当前课程资料中没有找到足够依据。\n"
        "不要编造课程资料中没有出现的内容。\n"
        "回答要求：\n"
        "1. 用中文回答。\n"
        "2. 面向初学者，解释要清楚。\n"
        "3. 可以适当分点说明。\n"
        "4. 如果涉及公式或概念，请先讲直观含义，再讲专业表达。\n"
        "5. 最后可以给出一句学习建议。\n"
        "6. 不要输出 JSON，只输出自然语言答案。"
    )


def _build_rag_user_prompt(question: str, context: str) -> str:
    """
    构造用户 prompt。
    """

    return (
        f"下面是从课程知识库中检索到的资料片段：\n\n"
        f"{context}\n\n"
        f"请你基于以上资料片段回答用户问题。\n\n"
        f"用户问题：{question}"
    )


def ask_knowledge_base(
    db: Session,
    course_id: str,
    question: str,
    top_k: int = 5,
    chapter_id: str | None = None,
    knowledge_point_id: str | None = None,
) -> dict:
    """
    RAG 知识库问答：
    1. 先向量检索
    2. 再把检索结果交给 DeepSeek 生成回答
    3. 返回 answer + references
    """

    if not question.strip():
        raise AppException(
            code=ErrorCode.PARAM_ERROR,
            message="问题不能为空",
            status_code=status.HTTP_400_BAD_REQUEST,
        )

    # 1. 先复用已有向量检索
    search_result = search_knowledge_chunks(
        db=db,
        course_id=course_id,
        query=question,
        top_k=top_k,
        chapter_id=chapter_id,
        knowledge_point_id=knowledge_point_id,
    )

    search_items = search_result.get("items", [])

    if not search_items:
        return {
            "course_id": course_id,
            "chapter_id": chapter_id,
            "knowledge_point_id": knowledge_point_id,
            "question": question,
            "answer": "当前课程资料中没有检索到足够相关的内容，暂时无法基于知识库回答这个问题。建议教师先上传相关课程资料，或扩大检索范围后重试。",
            "references": [],
            "llm_used": False,
            "provider": "DeepSeek",
        }

    # 2. 构造 references
    references = []

    for item in search_items:
        content = item.get("content") or ""
        preview = content[:180] + "..." if len(content) > 180 else content

        references.append(
            {
                "chunk_id": item.get("chunk_id"),
                "document_id": item.get("document_id"),
                "course_id": item.get("course_id"),
                "chapter_id": item.get("chapter_id"),
                "knowledge_point_id": item.get("knowledge_point_id"),
                "section": item.get("section"),
                "score": item.get("score", 0),
                "content_preview": preview,
            }
        )

    # 3. 构造 RAG prompt
    context = _build_rag_context(search_items)
    system_prompt = _build_rag_system_prompt()
    user_prompt = _build_rag_user_prompt(
        question=question,
        context=context,
    )

    # 4. 调用 llm
    try:
        llm = LLMService()

        # 这里我们需要一个普通文本生成方法
        answer = llm.rag_generate_text_sync(
            system_prompt=system_prompt,
            user_prompt=user_prompt,
            max_tokens=1200,
        )

        return {
            "course_id": course_id,
            "chapter_id": chapter_id,
            "knowledge_point_id": knowledge_point_id,
            "question": question,
            "answer": answer,
            "references": references,
            "llm_used": True,
            "provider": "llm",
        }

    except Exception as e:
        # LLM 失败时，不让整个接口崩，返回检索结果兜底
        fallback_answer = (
            "已检索到相关课程资料，但大模型生成回答失败。"
            "你可以先查看下方引用片段进行学习。\n\n"
            f"失败原因：{str(e)}"
        )

        return {
            "course_id": course_id,
            "chapter_id": chapter_id,
            "knowledge_point_id": knowledge_point_id,
            "question": question,
            "answer": fallback_answer,
            "references": references,
            "llm_used": False,
            "provider": "llm",
        }
